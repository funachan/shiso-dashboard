# -*- coding: utf-8 -*-
"""
宍粟市 人口ダッシュボード 自動更新スクリプト

宍粟市ホームページ「人口と世帯数」に掲載される Excel を取得し、
shiso_dashboard.html の以下を更新する。

  - MONTHLY_DATA : 町別の月次人口（まだ入っていない月をすべて追加）
  - AGE_DATA     : 最新月の年齢構成（0〜14歳 / 65歳以上）
  - ヘッダーの「基準日」「データ期間」表記
  - グラフ見出しの「（令和◯年◯月）」表記

【設計メモ】
市HPの Excel はファイル名の表記ゆれが激しい（gyouseikubetsujinko / gyoseikubetsujinko /
nenndobetsujinnko …）。月によっては「行政区別」と「5歳階級別」で名前が入れ替わっている
こともあるため、**ファイル名では判別せず、中身を見て判別する**。

  - ヘッダー行に「世帯数」がある          → 行政区別（町別の人口）
  - ヘッダー行に「0-4」が3ブロックある    → 5歳階級別（年齢構成）

列や行の位置も年度によってずれる可能性があるため、すべて見出し文字列を探して特定する。
"""

import io
import os
import re
import sys
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import load_workbook

# ------------------------------------------------------------------
# 設定
# ------------------------------------------------------------------
INDEX_URL = ("https://www.city.shiso.lg.jp/soshiki/shiminseikatsu/shimin/"
             "tantojoho/jinkoutokei/index.html")
HTML_FILE = "shiso_dashboard.html"

# 取りに行く年度ページの数（新しい方から）。年度またぎの取りこぼし防止に2つ見る
NENDO_PAGES = 2

TOWN_KEYS = [("山崎", "yamazaki"), ("一宮", "ichimiya"),
             ("波賀", "haga"), ("千種", "chigusa")]

HEADERS = {"User-Agent": "shiso-dashboard-updater (+https://github.com/funachan/shiso-dashboard)"}
TIMEOUT = 60


def log(msg):
    print(msg, flush=True)


def get(url, binary=False):
    r = requests.get(url, headers=HEADERS, timeout=TIMEOUT)
    r.raise_for_status()
    if binary:
        return r.content
    r.encoding = r.apparent_encoding or "utf-8"
    return r.text


# ------------------------------------------------------------------
# 1. 市HPから Excel のURLを集める
# ------------------------------------------------------------------
def collect_excel_urls():
    """{ 'YYYYMMDD': [url, ...] } を返す"""
    html = get(INDEX_URL)
    soup = BeautifulSoup(html, "html.parser")

    nendo_pages = []
    for a in soup.find_all("a"):
        text = a.get_text(strip=True)
        href = a.get("href")
        if href and "人口と世帯数" in text:
            nendo_pages.append(urljoin(INDEX_URL, href))
    if not nendo_pages:
        raise RuntimeError("年度別ページのリンクが見つかりませんでした")

    log(f"年度ページ {len(nendo_pages)} 件のうち新しい {NENDO_PAGES} 件を確認します")

    found = {}
    for page in nendo_pages[:NENDO_PAGES]:
        try:
            phtml = get(page)
        except Exception as e:
            log(f"  ! {page} の取得に失敗: {e}")
            continue
        psoup = BeautifulSoup(phtml, "html.parser")
        for a in psoup.find_all("a"):
            href = a.get("href")
            if not href or not href.lower().endswith(".xlsx"):
                continue
            m = re.search(r"(\d{8})\.xlsx$", href)
            if not m:
                continue
            date = m.group(1)
            url = urljoin(page, href)
            found.setdefault(date, [])
            if url not in found[date]:
                found[date].append(url)
    return found


# ------------------------------------------------------------------
# 2. Excel の中身を判別して読む
# ------------------------------------------------------------------
def sheet_rows(content):
    wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return rows


def _cell(v):
    return v.strip() if isinstance(v, str) else v


def _num(v):
    try:
        return int(float(v))
    except (TypeError, ValueError):
        return None


def parse_town_sheet(rows):
    """行政区別シート → {'total':n,'yamazaki':n,...} / 対象外なら None"""
    header_i = None
    for i, row in enumerate(rows[:20]):
        vals = [str(_cell(v)) for v in row if _cell(v) not in (None, "")]
        if any("世帯数" in v for v in vals) and any("合計" in v for v in vals):
            header_i = i
            break
    if header_i is None:
        return None

    header = [str(_cell(v)) if _cell(v) is not None else "" for v in rows[header_i]]
    total_col = None
    for j, h in enumerate(header):
        if h.replace(" ", "").replace("　", "") == "合計":
            total_col = j
    if total_col is None:
        return None

    wanted = {"合計": "total"}
    wanted.update({jp: en for jp, en in TOWN_KEYS})

    out = {}
    for row in rows[header_i + 1:]:
        if not row:
            continue
        label = _cell(row[0])
        if not isinstance(label, str):
            continue
        label = label.replace("町", "")
        if label in wanted and len(row) > total_col:
            n = _num(row[total_col])
            if n is not None:
                out[wanted[label]] = n

    if len(out) == 5:
        return out
    return None


def parse_age_sheet(rows):
    """5歳階級別シート → {'yamazaki':{'total','aged65','young014'}, ...} / 対象外なら None"""
    header_i = None
    for i, row in enumerate(rows[:20]):
        vals = [str(_cell(v)) for v in row]
        if vals.count("0-4") >= 3:
            header_i = i
            break
    if header_i is None:
        return None

    header = [str(_cell(v)) if _cell(v) is not None else "" for v in rows[header_i]]
    starts = [j for j, h in enumerate(header) if h == "0-4"]
    start = starts[2]                     # 3ブロック目＝男女計
    bands = header[start:start + 23]      # 0-4 … 110-
    if "65-69" not in bands:
        return None
    aged_from = start + bands.index("65-69")
    total_col = start + 23                # 直後が「合計」

    wanted = {"合計": "total"}
    wanted.update({jp: en for jp, en in TOWN_KEYS})

    out = {}
    for row in rows[header_i + 1:]:
        if not row:
            continue
        label = _cell(row[0])
        if not isinstance(label, str):
            continue
        label = label.replace("町", "")
        if label not in wanted or len(row) <= total_col:
            continue
        total = _num(row[total_col])
        if total is None:
            continue
        young = sum(_num(row[c]) or 0 for c in range(start, start + 3))
        aged = sum(_num(row[c]) or 0 for c in range(aged_from, start + 23))
        out[wanted[label]] = {"total": total, "aged65": aged, "young014": young}

    if len(out) == 5:
        return out
    return None


def fetch_month(urls):
    """その月のURL群から (town, age) を取り出す。取れなかった方は None"""
    town = age = None
    for url in urls:
        if town is not None and age is not None:
            break
        try:
            content = get(url, binary=True)
            rows = sheet_rows(content)
        except Exception as e:
            log(f"    ! {url.rsplit('/', 1)[-1]} を読めませんでした: {e}")
            continue
        if town is None:
            town = parse_town_sheet(rows)
            if town:
                log(f"    ✓ 行政区別: {url.rsplit('/', 1)[-1]}")
                continue
        if age is None:
            age = parse_age_sheet(rows)
            if age:
                log(f"    ✓ 5歳階級別: {url.rsplit('/', 1)[-1]}")
    return town, age


# ------------------------------------------------------------------
# 3. HTML の書き換え
# ------------------------------------------------------------------
def reiwa(year):
    return year - 2018


def existing_months(html):
    m = re.search(r"const MONTHLY_DATA = \[(.*?)\n\];", html, re.S)
    if not m:
        raise RuntimeError("MONTHLY_DATA が見つかりません")
    return set(re.findall(r'date:"(\d{4}-\d{2})"', m.group(1)))


def insert_monthly(html, rows_by_month):
    """rows_by_month: {'YYYY-MM': {'total':..,'yamazaki':..,...}}"""
    lines = []
    for ym in sorted(rows_by_month):
        d = rows_by_month[ym]
        lines.append(
            '  {{ date:"{ym}", total:{total}, yamazaki:{yamazaki}, '
            'ichimiya:{ichimiya}, haga:{haga}, chigusa:{chigusa} }},'.format(ym=ym, **d)
        )
    block = "\n".join(lines) + "\n"

    m = re.search(r"(const MONTHLY_DATA = \[.*?)(\n\];)", html, re.S)
    return html[:m.end(1)] + "\n" + block.rstrip("\n") + html[m.end(1):]


def insert_age(html, ym, age):
    if f'"{ym}": {{' in html:
        return html
    entry = (
        f'  "{ym}": {{\n'
        f'    yamazaki: {{ total:{age["yamazaki"]["total"]}, aged65:{age["yamazaki"]["aged65"]}, young014:{age["yamazaki"]["young014"]} }},\n'
        f'    ichimiya: {{ total:{age["ichimiya"]["total"]}, aged65:{age["ichimiya"]["aged65"]}, young014:{age["ichimiya"]["young014"]} }},\n'
        f'    haga:     {{ total:{age["haga"]["total"]}, aged65:{age["haga"]["aged65"]}, young014:{age["haga"]["young014"]} }},\n'
        f'    chigusa:  {{ total:{age["chigusa"]["total"]}, aged65:{age["chigusa"]["aged65"]}, young014:{age["chigusa"]["young014"]} }},\n'
        f'    total:    {{ total:{age["total"]["total"]}, aged65:{age["total"]["aged65"]}, young014:{age["total"]["young014"]} }},\n'
        f'  }},\n'
    )
    return html.replace("const AGE_DATA = {\n", "const AGE_DATA = {\n" + entry, 1)


def update_labels(html, date8):
    y, mth, day = int(date8[:4]), int(date8[4:6]), int(date8[6:])
    r = reiwa(y)
    html = re.sub(
        r"基準日：令和\d+年\d+月\d+日<br>データ期間：平成23年4月〜令和\d+年\d+月",
        f"基準日：令和{r}年{mth}月{day}日<br>データ期間：平成23年4月〜令和{r}年{mth}月",
        html)
    html = re.sub(r"（令和\d+年\d+月）", f"（令和{r}年{mth}月）", html)
    return html


# ------------------------------------------------------------------
# 4. メイン
# ------------------------------------------------------------------
def main():
    if not os.path.exists(HTML_FILE):
        log(f"❌ {HTML_FILE} が見つかりません")
        return 1

    html = io.open(HTML_FILE, encoding="utf-8").read()
    have = existing_months(html)
    log(f"現在の最新月: {max(have)}（全 {len(have)} か月）")

    available = collect_excel_urls()
    todo = sorted(d for d in available if f"{d[:4]}-{d[4:6]}" not in have)
    if not todo:
        log("✅ 追加すべき新しい月はありません（更新なし）")
        return 0

    log(f"追加対象: {', '.join(todo)}")

    new_rows = {}
    newest_age = None
    newest_date8 = None
    for date8 in todo:
        log(f"  {date8} を取得中…")
        town, age = fetch_month(available[date8])
        if not town:
            log(f"    ! {date8} の行政区別データが取れませんでした（スキップ）")
            continue
        new_rows[f"{date8[:4]}-{date8[4:6]}"] = town
        newest_date8 = date8
        if age:
            newest_age = age

    if not new_rows:
        log("❌ 新しいデータを1件も取得できませんでした")
        return 1

    html = insert_monthly(html, new_rows)
    if newest_age:
        html = insert_age(html, f"{newest_date8[:4]}-{newest_date8[4:6]}", newest_age)
    else:
        log("  ! 年齢構成データは取得できませんでした（人口のみ更新）")
    html = update_labels(html, newest_date8)

    io.open(HTML_FILE, "w", encoding="utf-8").write(html)
    log(f"✅ {HTML_FILE} を更新しました（{', '.join(sorted(new_rows))}）")

    summary = os.environ.get("GITHUB_STEP_SUMMARY")
    if summary:
        with io.open(summary, "a", encoding="utf-8") as f:
            f.write("## 人口データ更新完了\n\n")
            f.write("| 基準月 | 宍粟市計 | 山崎町 | 一宮町 | 波賀町 | 千種町 |\n")
            f.write("|---|---|---|---|---|---|\n")
            for ym in sorted(new_rows):
                d = new_rows[ym]
                f.write(f"| {ym} | {d['total']:,} | {d['yamazaki']:,} | "
                        f"{d['ichimiya']:,} | {d['haga']:,} | {d['chigusa']:,} |\n")
            if newest_age:
                t = newest_age["total"]
                rate = t["aged65"] / t["total"] * 100
                f.write(f"\n年齢構成も更新（高齢化率 {rate:.1f}%）\n")
    return 0


if __name__ == "__main__":
    sys.exit(main())
